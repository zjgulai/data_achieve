"""Regression coverage for the datasets workspace endpoints.

Self-seeding (unlike the shared `client` fixture) so it runs without a
pre-seeded demo workspace: it builds an in-memory sqlite DB, seeds a demo
workspace with one dataset, then exercises the list / preview / export /
download / job-lookup routes.
"""

from __future__ import annotations

import io
import uuid
from collections.abc import AsyncGenerator, AsyncIterator
from datetime import UTC, datetime

import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from data_intelligence_hub.core.config import get_settings
from data_intelligence_hub.core.database import get_session
from data_intelligence_hub.main import app
from data_intelligence_hub.models import (
    Base,
    CollectionTask,
    Dataset,
    DatasetVersion,
    Project,
    Source,
    TaskRun,
    User,
    Workspace,
)

NOW = datetime(2026, 10, 9, 12, 0, tzinfo=UTC)

USER_ID = uuid.uuid4()
WORKSPACE_ID = uuid.uuid4()
PROJECT_ID = uuid.uuid4()
SOURCE_ID = uuid.uuid4()
TASK_ID = uuid.uuid4()
RUN_ID = uuid.uuid4()
DATASET_ID = uuid.uuid4()
VERSION_ID = uuid.uuid4()

SELECTED_FIELDS = ["title", "price"]


async def _seed(session) -> None:
    session.add(
        User(
            id=USER_ID,
            email="owner@example.com",
            password_hash="hashed",
            name="Owner",
            status="active",
            created_at=NOW,
            updated_at=NOW,
        )
    )
    # Slug must match DEMO_WORKSPACE_SLUG so the list route resolves it.
    session.add(
        Workspace(
            id=WORKSPACE_ID,
            name="Demo",
            slug="data-achieve-demo",
            owner_id=USER_ID,
            created_at=NOW,
            updated_at=NOW,
        )
    )
    session.add(
        Project(
            id=PROJECT_ID,
            workspace_id=WORKSPACE_ID,
            name="采集项目",
            description="demo",
            domain="ecommerce",
            status="active",
            owner_id=USER_ID,
            created_at=NOW,
            updated_at=NOW,
        )
    )
    session.add(
        Source(
            id=SOURCE_ID,
            workspace_id=WORKSPACE_ID,
            project_id=PROJECT_ID,
            name="Amazon source",
            type="apify_actor",
            url=None,
            config={},
            schedule_cron=None,
            enabled=True,
            created_at=NOW,
            updated_at=NOW,
        )
    )
    # config.endpoint_type is what platform attribution walks.
    session.add(
        CollectionTask(
            id=TASK_ID,
            workspace_id=WORKSPACE_ID,
            project_id=PROJECT_ID,
            source_id=SOURCE_ID,
            collector_type="apify_actor",
            name="Amazon bestsellers",
            schedule_cron=None,
            status="enabled",
            config={"endpoint_type": "apify_amazon_bestsellers"},
            success_count=1,
            failure_count=0,
            last_run_at=NOW,
            created_at=NOW,
            updated_at=NOW,
        )
    )
    session.add(
        TaskRun(
            id=RUN_ID,
            task_id=TASK_ID,
            workspace_id=WORKSPACE_ID,
            status="success",
            started_at=NOW,
            finished_at=NOW,
            records_count=2,
            entities_count=0,
            error_message=None,
            error_traceback=None,
            logs=[],
            created_at=NOW,
        )
    )
    session.add(
        Dataset(
            id=DATASET_ID,
            workspace_id=WORKSPACE_ID,
            project_id=PROJECT_ID,
            name="Amazon 竞品数据集",
            dataset_type="ecommerce_product",
            status="active",
            description="amazon bestsellers sample",
            created_at=NOW,
            updated_at=NOW,
        )
    )
    rows = [
        {
            "row_id": f"row-{i}",
            "task_run_id": str(RUN_ID),
            "raw_record_id": str(uuid.uuid4()),
            "source_url": f"https://www.amazon.com/dp/B{i:08d}",
            "values": {"title": f"Product {i}", "price": 19.99 + i},
            "missing_fields": [],
            "completeness_percent": 100,
        }
        for i in range(2)
    ]
    session.add(
        DatasetVersion(
            id=VERSION_ID,
            dataset_id=DATASET_ID,
            workspace_id=WORKSPACE_ID,
            project_id=PROJECT_ID,
            created_by_user_id=USER_ID,
            cleaning_plan_id=None,
            source_workflow_run_id=None,
            # NB: the workflow-step/raw-record lineage columns are JSON, so a
            # Python None would be persisted as the JSON literal 'null' and
            # violate the lineage CHECK. Omit them to keep them SQL NULL.
            lineage_contract_version=None,
            version_number=1,
            source_task_run_ids=[str(RUN_ID)],
            selected_fields=SELECTED_FIELDS,
            cleaning_script=["trim string fields"],
            rows=rows,
            export_preview={"schema": {"title": "string", "price": "number"}},
            row_count=len(rows),
            average_completeness_percent=100,
            status="saved",
            created_at=NOW,
        )
    )
    await session.commit()


@pytest_asyncio.fixture()
async def api(tmp_path_factory) -> AsyncIterator[AsyncClient]:
    engine = create_async_engine(
        "sqlite+aiosqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)

    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with session_factory() as session:
        await _seed(session)

    async def override_session() -> AsyncGenerator[object, None]:
        async with session_factory() as session:
            yield session

    export_dir = tmp_path_factory.mktemp("dataset-exports")
    settings = get_settings()
    original_export_dir = settings.dataset_export_dir
    settings.dataset_export_dir = str(export_dir)

    app.dependency_overrides[get_session] = override_session
    transport = ASGITransport(app=app)
    try:
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            yield client
    finally:
        app.dependency_overrides.clear()
        settings.dataset_export_dir = original_export_dir
        await engine.dispose()


async def test_list_exposes_platform_category_and_timestamps(api: AsyncClient) -> None:
    response = await api.get("/api/automation/product-datasets")
    assert response.status_code == 200
    items = response.json()["items"]
    assert len(items) == 1
    item = items[0]
    assert item["dataset"]["dataset_type"] == "ecommerce_product"
    assert item["dataset"]["created_at"]
    assert item["dataset"]["updated_at"]
    # derived from config.endpoint_type -> catalog platform + content_type
    assert item["platforms"] == ["amazon"]
    assert item["category"] == "ecommerce"
    assert item["collector_types"] == ["apify_actor"]
    assert item["content_types"] == ["product"]


async def test_version_preview_returns_rows(api: AsyncClient) -> None:
    response = await api.get(
        f"/api/automation/product-datasets/{DATASET_ID}"
        f"/versions/{VERSION_ID}/preview?limit=1"
    )
    assert response.status_code == 200
    body = response.json()
    assert body["fields"] == SELECTED_FIELDS
    assert body["total_rows"] == 2
    assert body["preview_row_count"] == 1
    assert len(body["rows"]) == 1
    row = body["rows"][0]
    assert row["values"] == {"title": "Product 0", "price": 19.99}
    assert row["source_url"].startswith("https://www.amazon.com/dp/")


async def test_preview_404s_for_unknown_version(api: AsyncClient) -> None:
    response = await api.get(
        f"/api/automation/product-datasets/{DATASET_ID}"
        f"/versions/{uuid.uuid4()}/preview"
    )
    assert response.status_code == 404


async def test_export_requires_authorization_and_confirmation(api: AsyncClient) -> None:
    response = await api.post(
        "/api/automation/product-dataset-exports",
        json={
            "dataset_id": str(DATASET_ID),
            "dataset_version_id": str(VERSION_ID),
            "export_format": "csv",
        },
    )
    # authorized / confirm_create are required fields -> schema rejects
    assert response.status_code == 422


async def _create_export(api: AsyncClient, export_format: str) -> dict:
    response = await api.post(
        "/api/automation/product-dataset-exports",
        json={
            "authorized": True,
            "confirm_create": True,
            "dataset_id": str(DATASET_ID),
            "dataset_version_id": str(VERSION_ID),
            "export_format": export_format,
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


async def test_export_csv_is_downloadable(api: AsyncClient) -> None:
    job = await _create_export(api, "csv")
    # the response field the console must read is `id` (not `export_job_id`)
    assert job["id"]
    assert job["status"] == "success"
    assert job["download_url"]

    download = await api.get(job["download_url"])
    assert download.status_code == 200
    text = download.content.decode("utf-8")
    assert "title" in text and "Product 0" in text


async def test_get_export_job_by_id_route_exists(api: AsyncClient) -> None:
    job = await _create_export(api, "csv")
    response = await api.get(f"/api/automation/product-dataset-exports/{job['id']}")
    assert response.status_code == 200
    assert response.json()["id"] == job["id"]


async def test_get_export_job_unknown_id_is_404(api: AsyncClient) -> None:
    response = await api.get(f"/api/automation/product-dataset-exports/{uuid.uuid4()}")
    assert response.status_code == 404


async def test_export_xlsx_is_a_valid_workbook(api: AsyncClient) -> None:
    job = await _create_export(api, "xlsx")
    assert job["export_format"] == "xlsx"
    assert job["filename"].endswith(".xlsx")
    assert job["content_type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    download = await api.get(job["download_url"])
    assert download.status_code == 200
    assert download.content[:2] == b"PK"  # xlsx is a zip container

    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(download.content))
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    assert header[:2] == SELECTED_FIELDS
    assert sheet.max_row == 3  # header + 2 data rows


async def test_archive_hides_dataset_from_default_list(api: AsyncClient) -> None:
    """DELETE 走软删：默认列表消失，include_archived=true 仍可见。"""
    response = await api.delete(f"/api/automation/product-datasets/{DATASET_ID}")
    assert response.status_code == 200
    assert response.json()["status"] == "archived"

    listed = await api.get("/api/automation/product-datasets")
    assert listed.status_code == 200
    assert listed.json()["items"] == []
    assert listed.json()["total"] == 0

    archived_view = await api.get(
        "/api/automation/product-datasets?include_archived=true"
    )
    assert archived_view.status_code == 200
    assert len(archived_view.json()["items"]) == 1
    assert archived_view.json()["items"][0]["dataset"]["status"] == "archived"


async def test_archive_unknown_dataset_is_404(api: AsyncClient) -> None:
    response = await api.delete(
        f"/api/automation/product-datasets/{uuid.uuid4()}"
    )
    assert response.status_code == 404
    assert response.json()["detail"] == "dataset_not_found"


async def test_list_respects_limit_and_offset(api: AsyncClient) -> None:
    """offset/limit 分页切片 + total 用真实 count（不随页缩小）。"""
    full = await api.get("/api/automation/product-datasets?limit=100")
    assert full.status_code == 200
    assert full.json()["total"] == 1

    page = await api.get("/api/automation/product-datasets?limit=100&offset=1")
    assert page.status_code == 200
    assert page.json()["items"] == []
    assert page.json()["total"] == 1  # total 不随 offset 变化
